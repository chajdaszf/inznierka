#!/usr/bin/env python3
"""Generate real KiCad 9 sheets from the supplied workbook, with provenance.

Requires openpyxl and sexpdata==1.0.2. No PCB is generated. Sources from the
workbook are provisional until manufacturer and assembly checks are complete.
"""
import copy
import csv
import json
import math
import re
import uuid
from pathlib import Path

from openpyxl import load_workbook
import sexpdata as sx

ROOT = Path(__file__).resolve().parents[1]
NAME = 'Szafka_Glosowa_STM32_ESP32'
LIB = 'Szafka'
S = sx.Symbol
NAMESPACE = uuid.UUID('8b160458-58f8-4856-ac98-e7589c623d8c')


def uid(key):
    return str(uuid.uuid5(NAMESPACE, key))


def node(key, *args):
    return [S(key), *args]


def key(n):
    return str(n[0]) if isinstance(n, list) and n else ''


def get(n, tag):
    return next((a for a in n if key(a) == tag), None)


def walk(n, tag):
    if isinstance(n, list):
        if key(n) == tag:
            yield n
        else:
            for a in n:
                yield from walk(a, tag)


def effects(size=1, hide=False, justify=None):
    result = node('effects', node('font', node('size', size, size)))
    if justify:
        result.append(node('justify', S(justify)))
    if hide:
        result.append(node('hide', S('yes')))
    return result


def prop(name, val, x=0, y=0, hide=False, size=1, angle=0):
    return node('property', name, str(val), node('at', x, y, angle), effects(size, hide))


def save_sexpr(path, tree):
    # Individual top-level items are separate paragraphs for inspectability.
    text = '(' + str(tree[0]) + '\n'
    text += '\n'.join('  ' + sx.dumps(a) for a in tree[1:]) + '\n)\n'
    path.write_text(text, encoding='utf-8')


lib_cache = {}
symbols = {}
symbol_sources = {}


def system_symbol(lib, name):
    if lib not in lib_cache:
        source = Path('/usr/share/kicad/symbols') / (lib + '.kicad_sym')
        lib_cache[lib] = {a[1]: a for a in sx.load(source.open()) if key(a) == 'symbol'}
    original = lib_cache[lib][name]
    ext = get(original, 'extends')
    if ext:
        base = system_symbol(lib, ext[1])
        children = [copy.deepcopy(a) for a in base[2:] if key(a) == 'symbol']
        for a in children:
            a[1] = name + '_' + a[1].rsplit('_', 2)[-2] + '_' + a[1].rsplit('_', 1)[-1]
        result = [S('symbol'), name] + [copy.deepcopy(a) for a in original[2:] if key(a) != 'extends'] + children
    else:
        result = copy.deepcopy(original)
    return result


def use_symbol(lib, name):
    if name not in symbols:
        symbols[name] = system_symbol(lib, name)
        symbol_sources[name] = f'KiCad 9.0.2 {lib}:{name}; manufacturer recheck pending'
    return name


def make_drv():
    """Physical numbered pins 1..16 from uploaded map, NOT a verification claim.

    The exposed pad is named EP rather than guessing its footprint pad number.
    An exact RGT land pattern and EP mapping are intentionally not assigned.
    """
    name = 'DRV8876RGT_WORKBOOK'
    pins = [
        ('1', '~{SLEEP}', 'input', -17.78, 12.7, 0),
        ('2', '~{FAULT}', 'open_collector', -17.78, -12.7, 0),
        ('3', 'VREF', 'input', -17.78, 5.08, 0),
        ('4', 'IPROPI', 'output', 17.78, -10.16, 180),
        ('5', 'IMODE', 'input', -17.78, -7.62, 0),
        ('6', 'OUT1', 'output', 17.78, 5.08, 180),
        ('7', 'PGND', 'power_in', -2.54, -22.86, 90),
        ('8', 'OUT2', 'output', 17.78, 0, 180),
        ('9', 'VM', 'power_in', 0, 22.86, 270),
        ('10', 'VCP', 'passive', 17.78, 15.24, 180),
        ('11', 'CPH', 'passive', 17.78, 12.7, 180),
        ('12', 'CPL', 'passive', 17.78, 10.16, 180),
        ('13', 'GND', 'power_in', 0, -22.86, 90),
        ('14', 'PMODE', 'input', -17.78, -2.54, 0),
        ('15', 'EN/IN1', 'input', -17.78, 17.78, 0),
        ('16', 'PH/IN2', 'input', -17.78, 15.24, 0),
        ('EP', 'EP', 'power_in', 2.54, -22.86, 90),
    ]
    result = node('symbol', name, node('pin_names', node('offset', 1.016)),
                  node('in_bom', S('yes')), node('on_board', S('yes')),
                  prop('Reference', 'U'), prop('Value', 'DRV8876RGTR'),
                  prop('Footprint', '', hide=True),
                  prop('Datasheet', 'https://www.ti.com/lit/ds/symlink/drv8876.pdf', hide=True),
                  prop('Description', 'RGT pin map from workbook; EP/land pattern TO VERIFY', hide=True))
    body = node('symbol', name + '_0_1', node('rectangle', node('start', -12.7, 20.32),
                node('end', 12.7, -20.32), node('stroke', node('width', 0), node('type', S('default'))),
                node('fill', node('type', S('background')))))
    unit = node('symbol', name + '_1_1')
    for number, label, typ, x, y, angle in pins:
        unit.append(node('pin', S(typ), S('line'), node('at', x, y, angle),
                         node('length', 5.08 if abs(x) == 17.78 else 2.54),
                         node('name', label, effects(1.016)), node('number', number, effects(1.016))))
    result += [body, unit]
    symbols[name] = result
    symbol_sources[name] = 'Uploaded Mapa sygnalow ONLY; NOT independently verified; EP is a placeholder pad identifier'
    return name


workbook = load_workbook(ROOT / 'docs/input/BOM_i_polaczenia.xlsx', data_only=True)
parts = {}
input_rows = []
for row in workbook['BOM PCBA'].iter_rows(min_row=2, values_only=True):
    group, refs, quantity, mpn, description, lcsc, package, status, note, url = row
    input_rows.append(row)
    if refs.startswith('TP'):
        continue
    refs = refs.split(',')
    assert len(refs) == quantity, (refs, quantity)
    for ref in refs:
        assert ref not in parts, ref
        parts[ref] = {'MPN': mpn, 'LCSC': lcsc, 'description': description,
                      'note': note or '', 'url': url or '', 'package': package,
                      'Catalog_Status': 'FROM_WORKBOOK_NOT_INDEPENDENTLY_VERIFIED'}
parts['R18'] = copy.deepcopy(parts['R7'])
parts['R18']['note'] = 'Added explicit 10k nSLEEP pulldown; must check AND output loading'
parts['C28'] = copy.deepcopy(parts['C11'])
parts['C28']['note'] = 'Added 100nF local bypass for U8 AND gate'

root_id = uid('root')
sheet_names = ['Power_12V_5V_3V3', 'STM32_Control_SWD', 'ESP32_USB_Audio',
               'Motor_Drivers', 'Safety_Sensors_Connectors']
sheets = {}
pin_connections = []
placed_refs = set()


class Sheet:
    def __init__(self, name, page):
        self.name, self.page = name, page
        self.uuid = uid(name)
        self.items = []
        self.used = set()
        self.count = 0
        self.notes = []

    def text(self, msg, x, y, size=1.2):
        self.items.append(node('text', msg, node('at', x, y, 0), effects(size, justify='left'),
                               node('uuid', uid(self.name + ':text:' + str(len(self.items))))))

    def wire(self, x1, y1, x2, y2):
        if (x1, y1) == (x2, y2):
            return
        self.count += 1
        self.items.append(node('wire', node('pts', node('xy', x1, y1), node('xy', x2, y2)),
                               node('stroke', node('width', 0), node('type', S('default'))),
                               node('uuid', uid(self.name + ':wire:' + str(self.count)))))

    def label(self, net, x, y, angle=0):
        self.count += 1
        self.items.append(node('global_label', net, node('shape', S('bidirectional')),
                               node('at', x, y, angle),
                               effects(0.85, justify='right' if angle in (180, 270) else 'left'),
                               node('uuid', uid(self.name + ':label:' + str(self.count))),
                               prop('Intersheetrefs', '${INTERSHEET_REFS}', x, y, hide=True)))

    def part(self, ref, lib, name, x, y, nets, value=None, fp=None, rotation=0, note=''):
        if lib:
            name = use_symbol(lib, name)
        assert ref not in placed_refs, ref
        placed_refs.add(ref)
        self.used.add(name)
        definition = symbols[name]
        props = {p[1]: p[2] for p in definition if key(p) == 'property'}
        data = parts.get(ref, {})
        if value is None:
            value = data.get('MPN', props.get('Value', name))
        if fp is None:
            fp = props.get('Footprint', '')
        data['Footprint'] = fp
        data['Value'] = value
        data['Sheet'] = self.name
        if ref in parts:
            data['Footprint_Status'] = 'UNASSIGNED_PENDING_DATASHEET' if not fp else 'LIBRARY_CANDIDATE_PENDING_DATASHEET'
        pins = list(walk(definition, 'pin'))
        heights = [abs(get(p, 'at')[2]) for p in pins]
        # Long ICs get properties outside the symbol body, not over pin labels.
        tall = max(heights, default=3.81) > 15
        top = y - (max(heights) + 4.5 if tall else 6.5)
        if name.startswith('Conn_'):
            top = y - max(heights, default=0) - 7.62
        if name in ('USBLC6-2SC6', '74LVC1G08'):
            top = y - max(heights) - 12.7
        inst = node('symbol', node('lib_id', LIB + ':' + name), node('at', x, y, rotation),
                    node('unit', 1), node('in_bom', S('yes') if ref in parts else S('no')),
                    node('on_board', S('yes') if ref in parts else S('no')),
                    node('dnp', S('no')), node('uuid', uid('ref:' + ref)),
                    prop('Reference', ref, x, top, hide=ref.startswith('#'), angle=rotation % 180),
                    prop('Value', value, x, top + 2.8, size=1.016, angle=rotation % 180),
                    prop('Footprint', fp, x, y, hide=True),
                    prop('Datasheet', props.get('Datasheet', ''), x, y, hide=True))
        for field in ['LCSC', 'MPN', 'Catalog_Status', 'Footprint_Status']:
            if field in data:
                inst.append(prop(field, data[field], x, y, hide=True))
        if data or note:
            inst.append(prop('Design_Note', note or data.get('note', ''), x, y, hide=True))
        inst.append(node('instances', node('project', NAME, node('path', '/' + root_id + '/' + self.uuid,
                          node('reference', ref), node('unit', 1)))))
        self.items.append(inst)
        grouped = {}
        for p in pins:
            number, label = str(get(p, 'number')[1]), get(p, 'name')[1]
            a, b, angle = get(p, 'at')[1:]
            radians = math.radians(rotation)
            px = round(x + a * math.cos(radians) - b * math.sin(radians), 4)
            py = round(y - a * math.sin(radians) - b * math.cos(radians), 4)
            actual_angle = (angle + rotation) % 360
            assert number in nets, (ref, number, label)
            net = nets[number]
            inst.append(node('pin', number, node('uuid', uid(ref + ':pin:' + number))))
            pin_connections.append({'Sheet': self.name, 'Reference': ref, 'Pin': number,
                                    'Pin_Name': label, 'Net': net or 'NC',
                                    'Provenance': symbol_sources[name]})
            coord = (px, py, actual_angle)
            if coord in grouped:
                assert grouped[coord] == net, (ref, coord)
                continue
            grouped[coord] = net
            if net is None:
                self.items.append(node('no_connect', node('at', px, py),
                                       node('uuid', uid(ref + ':NC:' + number))))
            else:
                dx = round(-math.cos(math.radians(actual_angle)))
                dy = round(math.sin(math.radians(actual_angle)))
                ex, ey = round(px + dx * 5.08, 4), round(py + dy * 5.08, 4)
                self.wire(px, py, ex, ey)
                if dx:
                    self.label(net, ex, ey, 180 if dx < 0 else 0)
                else:
                    # A label at each vertical pin avoids connecting distinct power pins accidentally.
                    self.label(net, ex, ey, 270 if dy < 0 else 90)
        return data

    def pins_by_name(self, lib, name, assignments):
        use_symbol(lib, name)
        return {str(get(p, 'number')[1]): assignments.get(get(p, 'name')[1])
                for p in walk(symbols[name], 'pin')}

    def passive(self, ref, value, nets, x, y, kind='R', fp=None):
        if fp is None:
            fp = {'R': 'Resistor_SMD:R_0603_1608Metric', 'C': 'Capacitor_SMD:C_0603_1608Metric',
                  'LED': 'LED_SMD:LED_0603_1608Metric'}[kind]
        self.part(ref, 'Device', kind, x, y, {'1': nets[0], '2': nets[1]}, value, fp, rotation=90)

    def finish(self):
        embedded = []
        for name in sorted(self.used):
            sym = copy.deepcopy(symbols[name])
            sym[1] = LIB + ':' + name
            embedded.append(sym)
        tree = node('kicad_sch', node('version', 20250114), node('generator', S('eeschema')),
                    node('generator_version', '9.0'), node('uuid', self.uuid), node('paper', 'A3'),
                    node('title_block', node('title', NAME + ' / ' + self.name), node('rev', 'DRAFT 0.1'),
                         node('comment', 1, 'WORKBOOK INPUT; datasheet/catalog verification pending; NOT FOR ORDER')),
                    node('lib_symbols', *embedded), *self.items,
                    node('embedded_fonts', S('no')))
        save_sexpr(ROOT / (self.name + '.kicad_sch'), tree)


for index, sheetname in enumerate(sheet_names, 2):
    sheets[sheetname] = Sheet(sheetname, index)
    sheets[sheetname].text(sheetname, 15.24, 15.24, 2)
    sheets[sheetname].text('DRAFT - dane katalogowe i footprinty do potwierdzenia; NIE DO ZAMOWIENIA',
                           15.24, 20.32, 1)

# POWER: separate motor supply, buck supply and USB/logic rails.
p = sheets[sheet_names[0]]
p.part('J3', 'Connector_Generic', 'Conn_01x02', 40.64, 55.88,
       {'1': 'VIN_12V_EXTERNAL', '2': 'GND'}, fp='')
p.part('D1', 'Device', 'D_Schottky', 104.14, 55.88,
       {'1': 'VM_MOTORS', '2': 'VIN_12V_EXTERNAL'}, value='SS54', fp='Diode_SMD:D_SMB', rotation=180)
p.part('D4', 'Device', 'D_TVS', 175.26, 55.88,
       {'1': 'GND', '2': 'VM_MOTORS'}, value='SMBJ15A', fp='Diode_SMD:D_SMB')
p.part('C1', 'Device', 'C_Polarized', 246.38, 55.88,
       {'1': 'VM_MOTORS', '2': 'GND'}, value='100uF 25V', fp='')
p.text('Zewnetrzny bezpiecznik DC TBD przed J3. Na PCB tylko 12V DC.', 15.24, 81.28)
p.text('D1 straty/termika i D4 koordynacja z bezpiecznikiem: TO VERIFY.', 15.24, 88.9)
for index, (ref, raw, out, sw, bst, inductor, lvalue, lfp) in enumerate([
    ('U5', 'VM_MOTORS', '5V_BUCK', 'BUCK5_SW', 'BUCK5_BST', 'L1', '4.7uH', ''),
    ('U6', '5V_OR', '+3V3', 'BUCK3_SW', 'BUCK3_BST', 'L2', '3.9uH', 'Inductor_SMD:L_Bourns-SRN6028')]):
    y = 119.38 + index * 63.5
    p.part(ref, 'Regulator_Switching', 'AP63205WU' if ref == 'U5' else 'AP63203WU', 81.28, y,
           {'1': out, '2': raw, '3': raw, '4': 'GND', '5': sw, '6': bst}, value=parts[ref]['MPN'])
    p.part(inductor, 'Device', 'L', 160.02, y, {'1': sw, '2': out}, value=lvalue, fp=lfp, rotation=90)
    base = 2 if index == 0 else 6
    p.passive('C' + str(base), '10uF 25V', (raw, 'GND'), 243.84, y-15.24,
              kind='C', fp='Capacitor_SMD:C_0805_2012Metric')
    p.passive('C' + str(base + 1), '100nF 50V', (bst, sw), 243.84, y+10.16, kind='C')
    for col, cr in enumerate([base + 2, base + 3]):
        p.passive('C' + str(cr), '22uF 25V', (out, 'GND'), 337.82, y-15.24+col*25.4,
                  kind='C', fp='Capacitor_SMD:C_0805_2012Metric')
for ref, inputnet, x in [('D2', '5V_BUCK', 60.96), ('D3', 'USB_VBUS', 157.48)]:
    p.part(ref, 'Device', 'D_Schottky', x, 236.22,
           {'1': '5V_OR', '2': inputnet}, value='SS14', fp='Diode_SMD:D_SMA', rotation=180)
p.text('USB -> D3 -> 5V_OR -> U6 -> +3V3. USB nie ma drogi zasilania VM_MOTORS.', 15.24, 256.54)
p.text('U5/U6: stale napiecia; FB do wyjscia ZA dlawikiem. CIN/COUT i BST blisko ukladow.', 15.24, 264.16)
for i, net in enumerate(['VIN_12V_EXTERNAL', 'VM_MOTORS', 'USB_VBUS', '5V_BUCK', '5V_OR', '+3V3', 'GND']):
    p.part('#FLG0' + str(i+1), 'power', 'PWR_FLAG', 302.26+i*12.7, 241.3,
           {'1': net}, value='PWR_FLAG', fp='')

# STM: physical pins are from the installed STM32G030C8Tx symbol.
p = sheets[sheet_names[1]]
stm = {
    'NRST': 'STM_NRST', 'VREF+': '+3V3', 'VBAT': '+3V3', 'VDD': '+3V3', 'VSS': 'GND',
    'PA0': 'M1_CURRENT_ADC', 'PA1': 'M2_CURRENT_ADC', 'PA6': 'M1_IN1', 'PA7': 'M1_IN2',
    'PB0': 'M2_IN1', 'PB1': 'M2_IN2', 'PA8': 'MOTOR_MCU_ENABLE', 'PB3': 'M1_nFAULT', 'PB4': 'M2_nFAULT',
    'PA9/PA11': 'STM_TX_ESP_RX', 'PA10/PA12': 'ESP_TX_STM_RX', 'PB6': 'I2C_SCL', 'PB7': 'I2C_SDA',
    'PB8': 'TOF1_XSHUT', 'PB9': 'TOF2_XSHUT', 'PB10': 'LIM_1_OPEN_NC', 'PB11': 'LIM_1_CLOSED_NC',
    'PB12': 'LIM_2_OPEN_NC', 'PB13': 'LIM_2_CLOSED_NC', 'PB14': 'SAFETY_LOOP_OK',
    'PB15': 'MCU_STATUS_LED', 'PA13': 'SWDIO', 'PA14': 'SWCLK'
}
p.part('U1', 'MCU_ST_STM32G0', 'STM32G030C8Tx', 91.44, 119.38,
       p.pins_by_name('MCU_ST_STM32G0', 'STM32G030C8Tx', stm), value='STM32G030C8T6')
p.part('J2', 'Connector_Generic', 'Conn_02x07_Odd_Even', 246.38, 60.96,
       {str(n): {3: '+3V3', 4: 'SWDIO', 5: 'GND', 6: 'SWCLK', 7: 'GND',
                 11: 'GND', 12: 'STM_NRST'}.get(n) for n in range(1,15)}, fp='')
p.text('STDC14 wg arkusza: 3 VTref, 4 SWDIO, 6 SWCLK, 12 NRST.', 185.42, 86.36)
p.text('1/2/9 reserved; 8/10/13/14 nieuzywane. J2 footprint do weryfikacji.', 185.42, 93.98)
for ref, net, x in [('C11', '+3V3', 203.2), ('C12', '+3V3', 279.4), ('C13', '+3V3', 355.6)]:
    p.passive(ref, '100nF 50V', (net, 'GND'), x, 127, kind='C')
p.text('C11 VDD, C12 VBAT, C13 VREF+; blisko odpowiednich pinow U1.', 185.42, 147.32)
p.passive('R9', '10k', ('+3V3', 'STM_NRST'), 223.52, 172.72)
p.passive('C16', '100nF 50V', ('STM_NRST', 'GND'), 330.2, 172.72, kind='C')
p.passive('R22', '1k', ('+3V3', 'LED_POWER_A'), 203.2, 213.36)
p.passive('LED1', 'RED', ('GND', 'LED_POWER_A'), 279.4, 213.36, kind='LED')
p.passive('R23', '1k', ('MCU_STATUS_LED', 'LED_STATUS_A'), 203.2, 248.92)
p.passive('LED2', 'RED', ('GND', 'LED_STATUS_A'), 279.4, 248.92, kind='LED')
p.text('UART: PA9/PA11 = pin33, PA10/PA12 = pin34; piny29/32 NC.', 15.24, 203.2)
p.text('Sprawdz remap PA9/PA10 i AF w datasheet/reference manual przed firmware.', 15.24, 210.82)
p.text('HSI bez kwarcu. Nieuzywane GPIO oznaczone NC; krańcówki wymagaja pull-up w firmware.', 15.24, 226.06, 1)

# ESP / native USB / microphone.
p = sheets[sheet_names[2]]
esp = {'3V3': '+3V3', 'GND': 'GND', 'EN': 'ESP_EN', 'IO0': 'ESP_BOOT',
       'IO4': 'I2S_BCLK', 'IO5': 'I2S_WS', 'IO6': 'I2S_DATA', 'IO17': 'ESP_TX_STM_RX',
       'IO18': 'STM_TX_ESP_RX', 'USB_D-': 'USB_D_MINUS_MCU', 'USB_D+': 'USB_D_PLUS_MCU'}
p.part('U2', 'RF_Module', 'ESP32-S3-WROOM-1', 86.36, 106.68,
       p.pins_by_name('RF_Module', 'ESP32-S3-WROOM-1', esp), value='ESP32-S3-WROOM-1-N8')
usb = {'S1': 'GND', **{x:'GND' for x in ['A1','A12','B1','B12']},
       **{x:'USB_VBUS' for x in ['A4','A9','B4','B9']}, 'A5':'USB_CC1', 'B5':'USB_CC2',
       'A7':'USB_D_MINUS_CONN', 'B7':'USB_D_MINUS_CONN', 'A6':'USB_D_PLUS_CONN',
       'B6':'USB_D_PLUS_CONN', 'A8':None, 'B8':None}
p.part('J1', 'Connector', 'USB_C_Receptacle_USB2.0_16P', 241.3, 66.04, usb,
       value='TYPE-C-31-M-12', fp='Connector_USB:USB_C_Receptacle_HRO_TYPE-C-31-M-12')
p.part('U7', 'Power_Protection', 'USBLC6-2SC6', 241.3, 124.46,
       {'1':'USB_D_MINUS_CONN','6':'USB_D_MINUS_CONN','3':'USB_D_PLUS_CONN',
        '4':'USB_D_PLUS_CONN','5':'USB_VBUS','2':'GND'}, value='USBLC6-2SC6')
for ref, val, nets, x, y in [
    ('R1','5.1k',('USB_CC1','GND'),345.44,60.96),
    ('R2','5.1k',('USB_CC2','GND'),345.44,91.44),
    ('R19','22R',('USB_D_MINUS_CONN','USB_D_MINUS_MCU'),223.52,167.64),
    ('R20','22R',('USB_D_PLUS_CONN','USB_D_PLUS_MCU'),223.52,200.66),
    ('R3','10k',('+3V3','ESP_EN'),55.88,172.72),
    ('R4','10k',('+3V3','ESP_BOOT'),137.16,172.72)]:
    p.passive(ref,val,nets,x,y)
for ref, net, x in [('SW1','ESP_BOOT',137.16),('SW2','ESP_EN',55.88)]:
    p.part(ref,'Switch','SW_Push',x,205.74,{'1':net,'2':'GND'},value=parts[ref]['MPN'],
           fp='Button_Switch_SMD:SW_Push_1P1T_XKB_TS-1187A')
for ref,val,nets,x,y,fp in [
    ('C10','22uF 25V',('+3V3','GND'),55.88,241.3,'Capacitor_SMD:C_0805_2012Metric'),
    ('C14','100nF 50V',('+3V3','GND'),137.16,241.3,None),
    ('C15','1uF 50V',('ESP_EN','GND'),55.88,269.24,None)]:
    p.passive(ref,val,nets,x,y,kind='C',fp=fp)
p.part('J8','Connector_Generic','Conn_01x06',345.44,170.18,
       {'1':'+3V3','2':'GND','3':'I2S_BCLK','4':'I2S_WS','5':'I2S_DATA','6':'GND'},
       fp='Connector_JST:JST_PH_S6B-PH-SM4-TB_1x06-1MP_P2.00mm_Horizontal')
p.text('INMP441 zewnetrzny: pin6 L/R = GND, lewy kanal. Potwierdz kabel/modul.',294.64,203.2,1)
p.text('USB GPIO19/20 = piny modulu13/14; UART GPIO17/18 = 10/11.',185.42,233.68)
p.text('J1: sygnaly SMT, kotwy oslony THT! PCBA i wymaganie SMD do uzgodnienia.',185.42,246.38,1)
p.text('22R wstepne; antena RF keepout, impedancja USB i ESD przy J1 do layout review.',185.42,259.08,1)

# MOTOR: physical numbered pins from workbook, with explicit verification hold.
p = sheets[sheet_names[3]]
drv_name = make_drv()
for m, x in [(1,76.2),(2,279.4)]:
    ref='U'+str(m+2)
    nets={'1':'MOTOR_DRV_ENABLE','2':f'M{m}_nFAULT','3':f'M{m}_VREF','4':f'M{m}_CURRENT_ADC',
          '5':None,'6':f'MOTOR_{m}_A','7':'GND','8':f'MOTOR_{m}_B','9':'VM_MOTORS',
          '10':f'M{m}_VCP','11':f'M{m}_CPH','12':f'M{m}_CPL','13':'GND','14':'+3V3',
          '15':f'M{m}_IN1','16':f'M{m}_IN2','EP':'GND'}
    p.part(ref,None,drv_name,x,91.44,nets,value='DRV8876RGTR',fp='')
    p.text('RGT z mapy XLSX; pin EP i land pattern NIEPOTWIERDZONE.',x-55.88,139.7,1)
    p.text('IMODE Hi-Z zamierzone wg XLSX; PMODE high (PWM).',x-55.88,147.32,1)
    # Protective and configuration passives below each bridge.
    entries=[
      ('R'+str(m+4),'10k',('+3V3',f'M{m}_nFAULT'),'R',None),
      ('R'+str(m+11),'2.49k',(f'M{m}_CURRENT_ADC','GND'),'R',None),
      ('R'+str(m+15),'6.8k',('+3V3',f'M{m}_VREF'),'R',None),
      ('R'+str(m+13),'10k',(f'M{m}_VREF','GND'),'R',None),
      ('C'+str(m+16),'100nF 50V',('VM_MOTORS','GND'),'C',None),
      ('C'+str(m+20),'100nF 50V',(f'M{m}_VCP','VM_MOTORS'),'C',None),
      ('C'+str(m+22),'22nF 50V',(f'M{m}_CPH',f'M{m}_CPL'),'C',None),
      ('C'+str(m+25),'1nF 50V',(f'M{m}_CURRENT_ADC','GND'),'C',None),
    ]
    for j,(r,v,ns,typ,fp) in enumerate(entries):
        px=x-30.48+(j%2)*88.9
        py=172.72+(j//2)*25.4
        p.passive(r,v,ns,px,py,kind=typ,fp=fp)
    p.part('C'+str(m+18),'Device','C_Polarized',x+58.42,119.38,
           {'1':'VM_MOTORS','2':'GND'},value='100uF 25V',fp='')
    p.part('J'+str(m+3),'Connector_Generic','Conn_01x02',x+60.96,55.88,
           {'1':f'MOTOR_{m}_A','2':f'MOTOR_{m}_B'},fp='')
p.text('VREF = 3.3V * 10k/(6.8k+10k) = 1.964V. WARTOSC TYLKO WSTEPNA.',15.24,266.7,1.1)
p.text('RIPROPI/VREF i wspolczynnik sense musza byc sprawdzone dla wybranego silnika.',15.24,274.32,1)

# SAFETY, switches and ToF: physical connector numbering from workbook.
p = sheets[sheet_names[4]]
for ref,nets,x in [('J13',('+3V3','SAFETY_LINK_1'),50.8),
                    ('J14',('SAFETY_LINK_1','SAFETY_LINK_2'),162.56),
                    ('J15',('SAFETY_LINK_2','SAFETY_LOOP_OK'),274.32)]:
    p.part(ref,'Connector_Generic','Conn_01x02',x,55.88,{'1':nets[0],'2':nets[1]},
           fp='Connector_JST:JST_PH_S2B-PH-SM4-TB_1x02-1MP_P2.00mm_Horizontal')
p.text('J13 E-STOP NC -> J14 kontakt szuflady1 NC -> J15 kontakt szuflady2 NC.',15.24,81.28)
p.text('Rozwarcie styku obniza SAFETY_LOOP_OK przez R8, niezaleznie od kodu STM.',15.24,88.9)
p.part('U8','74xGxx','74LVC1G08',81.28,129.54,
       {'1':'MOTOR_MCU_ENABLE','2':'SAFETY_LOOP_OK','3':'GND','4':'MOTOR_DRV_ENABLE','5':'+3V3'},
       value='SN74LVC1G08DBVR',fp='Package_TO_SOT_SMD:SOT-23-5')
for ref,val,nets,x,y,typ in [
    ('R7','10k',('MOTOR_MCU_ENABLE','GND'),187.96,106.68,'R'),
    ('R8','10k',('SAFETY_LOOP_OK','GND'),294.64,106.68,'R'),
    ('R18','10k',('MOTOR_DRV_ENABLE','GND'),187.96,137.16,'R'),
    ('C25','1nF 50V',('SAFETY_LOOP_OK','GND'),294.64,137.16,'C'),
    ('C28','100nF 50V',('+3V3','GND'),81.28,170.18,'C')]:
    p.passive(ref,val,nets,x,y,kind=typ)
p.text('R7 + R8 default OFF; R18 dodatkowy pulldown nSLEEP; C28 bypass U8.',15.24,187.96)
for ref,net,x in [('J9','LIM_1_OPEN_NC',45.72),('J10','LIM_1_CLOSED_NC',139.7),
                  ('J11','LIM_2_OPEN_NC',233.68),('J12','LIM_2_CLOSED_NC',327.66)]:
    p.part(ref,'Connector_Generic','Conn_01x02',x,218.44,{'1':net,'2':'GND'},
           fp='Connector_JST:JST_PH_S2B-PH-SM4-TB_1x02-1MP_P2.00mm_Horizontal')
for m,x in [(1,63.5),(2,187.96)]:
    p.part('J'+str(m+5),'Connector_Generic','Conn_01x05',x,261.62,
           {'1':'+3V3','2':'GND','3':'I2C_SCL','4':'I2C_SDA','5':f'TOF{m}_XSHUT'},
           fp='Connector_JST:JST_PH_S5B-PH-SM4-TB_1x05-1MP_P2.00mm_Horizontal')
p.passive('R10','4.7k',('+3V3','I2C_SCL'),355.6,165.1)
p.passive('R11','4.7k',('+3V3','I2C_SDA'),355.6,190.5)
p.text('Prototyp NC/ToF/IPROPI NIE jest certyfikowanym ukladem ochrony BHP.',15.24,198.12,1.3)

# Explicit power symbols, separate from nonphysical PWR_FLAGs.
for i,p in enumerate(sheets.values(),1):
    p.part(f'#PWR0{i}1','power','GND',393.7,35.56,{'1':'GND'},value='GND',fp='')
    p.part(f'#PWR0{i}2','power','+3V3',368.3,35.56,{'1':'+3V3'},value='+3V3',fp='')
    p.finish()

assert set(parts) == {r for r in placed_refs if not r.startswith('#')}, (set(parts)-placed_refs,placed_refs-set(parts))

root = node('kicad_sch',node('version',20250114),node('generator',S('eeschema')),
            node('generator_version','9.0'),node('uuid',root_id),node('paper','A3'),
            node('title_block',node('title',NAME),node('rev','DRAFT 0.1'),
                 node('comment',1,'NOT FOR ORDER - manufacturer/catalog/actuator verification pending')),
            node('lib_symbols'))
root += [node('text','INTELIGENTNA SZAFKA - STM32 + ESP32 / 2 szuflady',node('at',25.4,22.86,0),
              effects(2,justify='left'),node('uuid',uid('root:title')))]
for i,name in enumerate(sheet_names):
    x,y=25.4+(i%2)*187.96,43.18+(i//2)*71.12
    sh=sheets[name]
    root.append(node('sheet',node('at',x,y),node('size',154.94,50.8),
                     node('stroke',node('width',0.254),node('type',S('default'))),
                     node('fill',node('color',0,0,0,0)),node('uuid',sh.uuid),
                     prop('Sheetname',name,x+77.47,y-3.81),
                     prop('Sheetfile',name+'.kicad_sch',x+77.47,y+54.61),
                     node('instances',node('project',NAME,node('path','/'+root_id,node('page',str(sh.page)))))))
    summaries = [
        'Wejscie 12V DC / SS54 / TVS\nAP63205: 12V -> 5V_BUCK\nSS14 OR: USB i buck -> 5V_OR\nAP63203: 5V_OR -> +3V3',
        'STM32G030C8T6 / LQFP48\nSWD STDC14 / UART do ESP32\nPWM napedow / ADC IPROPI\nI2C ToF / krańcówki / status',
        'ESP32-S3-WROOM-1-N8\nNatywne USB-C / ESD / 22R\nPrzyciski BOOT i RESET\nI2S do zewnetrznego INMP441',
        '2x DRV8876RGTR (PINOUT TO VERIFY)\nPWM / IPROPI / VREF wstepne\nBramka AND steruje nSLEEP\nDwa niezalezne wyjscia silnikowe',
        'NC: E-STOP + kontakty obu szuflad\nAND: MCU_ENABLE && SAFETY_OK\n4 osobne krańcówki NC\n2x ToF: I2C + oddzielne XSHUT'
    ]
    root.append(node('text',summaries[i],node('at',x+7.62,y+10.16,0),effects(1.4,justify='left'),
                     node('uuid',uid('root:summary:'+name))))
root += [node('text','Arkusze lacza etykiety globalne. Zasilanie: 12V -> VM_MOTORS; USB tylko 5V_OR / +3V3.',
              node('at',25.4,259.08,0),effects(1.2,justify='left'),node('uuid',uid('root:power'))),
         node('text','SCHEMAT ROBOCZY. Brak weryfikacji katalogowej nie jest wynikiem ERC. Bez PCB/Gerber/CPL.',
              node('at',25.4,269.24,0),effects(1.2,justify='left'),node('uuid',uid('root:hold'))),
         node('sheet_instances',node('path','/',node('page','1'))),node('embedded_fonts',S('no'))]
save_sexpr(ROOT/(NAME+'.kicad_sch'),root)

project={'meta':{'filename':NAME+'.kicad_pro','version':1},'schematic':{'meta':{'version':1}},
         'erc':{'meta':{'version':0},'erc_exclusions':[]},'text_variables':{}}
(ROOT/(NAME+'.kicad_pro')).write_text(json.dumps(project,indent=2)+'\n')
save_sexpr(ROOT/(LIB+'.kicad_sym'),node('kicad_symbol_lib',node('version',20241209),
            node('generator',S('kicad_symbol_editor')),node('generator_version','9.0'),
            *[symbols[n] for n in sorted(symbols)]))
save_sexpr(ROOT/'sym-lib-table',node('sym_lib_table',node('version',7),
            node('lib',node('name',LIB),node('type','KiCad'),node('uri','${KIPRJMOD}/'+LIB+'.kicad_sym'),
                 node('options',''),node('descr','Pinned KiCad symbols plus provisional workbook DRV8876 RGT map'))))

def csv_out(path,rows,fields):
    with path.open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=fields)
        w.writeheader()
        w.writerows(rows)

csv_out(ROOT/'reports/PIN_NET_MAP.csv',pin_connections,
        ['Sheet','Reference','Pin','Pin_Name','Net','Provenance'])
csv_out(ROOT/'reports/PART_VERIFICATION.csv',
        [{'Reference':r,**{f:d.get(f,'') for f in ['MPN','LCSC','Value','Footprint','Catalog_Status','Footprint_Status','url','note']}}
         for r,d in parts.items()],
        ['Reference','MPN','LCSC','Value','Footprint','Catalog_Status','Footprint_Status','url','note'])
(ROOT/'reports/expected_connectivity.json').write_text(json.dumps(pin_connections,indent=2,ensure_ascii=False)+'\n')
print(f'Generated {NAME}: 5 child sheets, {len(parts)} physical parts, {len(pin_connections)} pin records.')
